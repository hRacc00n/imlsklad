import os
import json
from flask import request, jsonify, send_file
from datetime import datetime, timedelta
from db.database import get_db
from db.models import Order, PersonalTask, OrderError, ErrorType, Complaint, ComplaintCategory
from utils.file_loader import load_json
import io

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False


def register_statistics_routes(app):
    
    # ===== GET: Основные показатели за период =====
    @app.route('/api/statistics', methods=['GET'])
    def get_statistics():
        date_from_str = request.args.get('date_from', '')
        date_to_str = request.args.get('date_to', '')
        
        # Парсим даты
        try:
            if date_from_str:
                date_from = datetime.fromisoformat(date_from_str.replace('T', ' '))
            else:
                date_from = datetime.utcnow().replace(day=1, hour=0, minute=0, second=0, microsecond=0)
            
            if date_to_str:
                date_to = datetime.fromisoformat(date_to_str.replace('T', ' '))
            else:
                date_to = datetime.utcnow()
        except:
            return jsonify({'error': 'Неверный формат даты'}), 400
        
        # Устанавливаем конец дня для date_to
        date_to = date_to.replace(hour=23, minute=59, second=59)
        
        with get_db() as db:
            # ===== 1. Отгрузки выполненные (regions + spb) =====
            shipments_completed = db.query(Order).filter(
                Order.type.in_(['regions', 'spb']),
                Order.status == 'completed',
                Order.completed_at >= date_from,
                Order.completed_at <= date_to
            ).count()
            
            # ===== 2. Отгрузки с ошибками =====
            shipments_with_errors = db.query(OrderError).join(
                Order, Order.id == OrderError.order_id
            ).filter(
                Order.type.in_(['regions', 'spb']),
                OrderError.created_at >= date_from,
                OrderError.created_at <= date_to
            ).distinct(OrderError.order_id).count()
            
            # ===== 3. Задачи поставлено (все хабы + личные) =====
            orders_created = db.query(Order).filter(
                Order.created_at >= date_from,
                Order.created_at <= date_to
            ).count()
            
            personal_tasks_created = db.query(PersonalTask).filter(
                PersonalTask.created_at >= date_from,
                PersonalTask.created_at <= date_to
            ).count()
            
            tasks_created = orders_created + personal_tasks_created
            
            # ===== 4. Задачи выполнено (все хабы + личные) =====
            orders_completed = db.query(Order).filter(
                Order.status == 'completed',
                Order.completed_at >= date_from,
                Order.completed_at <= date_to
            ).count()
            
            personal_tasks_completed = db.query(PersonalTask).filter(
                PersonalTask.status == 'completed',
                PersonalTask.completed_at >= date_from,
                PersonalTask.completed_at <= date_to
            ).count()
            
            tasks_completed = orders_completed + personal_tasks_completed
            
            # ===== 5. Жалобы подтверждённые =====
            complaints_confirmed = db.query(Complaint).filter(
                Complaint.status == 'confirmed',
                Complaint.created_at >= date_from,
                Complaint.created_at <= date_to
            ).count()
            
            return jsonify({
                'period': {
                    'date_from': date_from.strftime('%Y-%m-%d'),
                    'date_to': date_to.strftime('%Y-%m-%d')
                },
                'indicators': {
                    'shipments_completed': shipments_completed,
                    'shipments_with_errors': shipments_with_errors,
                    'tasks_created': tasks_created,
                    'tasks_completed': tasks_completed,
                    'complaints_confirmed': complaints_confirmed
                },
                'details': {
                    'orders_created': orders_created,
                    'personal_tasks_created': personal_tasks_created,
                    'orders_completed': orders_completed,
                    'personal_tasks_completed': personal_tasks_completed
                }
            }), 200
    
    # ===== GET: Список ошибок за период =====
    @app.route('/api/statistics/errors', methods=['GET'])
    def get_statistics_errors():
        date_from_str = request.args.get('date_from', '')
        date_to_str = request.args.get('date_to', '')
        
        try:
            if date_from_str:
                date_from = datetime.fromisoformat(date_from_str.replace('T', ' '))
            else:
                date_from = datetime.utcnow().replace(day=1, hour=0, minute=0, second=0, microsecond=0)
            
            if date_to_str:
                date_to = datetime.fromisoformat(date_to_str.replace('T', ' '))
            else:
                date_to = datetime.utcnow()
        except:
            return jsonify({'error': 'Неверный формат даты'}), 400
        
        date_to = date_to.replace(hour=23, minute=59, second=59)
        
        with get_db() as db:
            errors = db.query(OrderError).filter(
                OrderError.created_at >= date_from,
                OrderError.created_at <= date_to
            ).order_by(OrderError.created_at.desc()).all()
            
            result = []
            for error in errors:
                error_dict = error.to_dict()
                
                # Тип ошибки
                error_type = db.query(ErrorType).filter(ErrorType.id == error.error_type_id).first()
                if error_type:
                    error_dict['error_type_name'] = error_type.name
                
                # Информация о заказе
                order = db.query(Order).filter(Order.id == error.order_id).first()
                if order:
                    email_data = {}
                    if order.email_data:
                        try:
                            email_data = json.loads(order.email_data)
                        except:
                            pass
                    
                    error_dict['order_number'] = email_data.get('order_number', '')
                    error_dict['contractor'] = email_data.get('contractor', '')
                    error_dict['order_type'] = order.type
                
                result.append(error_dict)
            
            return jsonify(result), 200
    
    # ===== GET: Список жалоб за период =====
    @app.route('/api/statistics/complaints', methods=['GET'])
    def get_statistics_complaints():
        date_from_str = request.args.get('date_from', '')
        date_to_str = request.args.get('date_to', '')
        
        try:
            if date_from_str:
                date_from = datetime.fromisoformat(date_from_str.replace('T', ' '))
            else:
                date_from = datetime.utcnow().replace(day=1, hour=0, minute=0, second=0, microsecond=0)
            
            if date_to_str:
                date_to = datetime.fromisoformat(date_to_str.replace('T', ' '))
            else:
                date_to = datetime.utcnow()
        except:
            return jsonify({'error': 'Неверный формат даты'}), 400
        
        date_to = date_to.replace(hour=23, minute=59, second=59)
        
        with get_db() as db:
            complaints = db.query(Complaint).filter(
                Complaint.created_at >= date_from,
                Complaint.created_at <= date_to
            ).order_by(Complaint.created_at.desc()).all()
            
            result = []
            for complaint in complaints:
                complaint_dict = complaint.to_dict()
                
                # Категория
                category = db.query(ComplaintCategory).filter(
                    ComplaintCategory.id == complaint.category_id
                ).first()
                if category:
                    complaint_dict['category_name'] = category.name
                
                result.append(complaint_dict)
            
            return jsonify(result), 200
    
    # ===== GET: Загрузка менеджеров =====
    @app.route('/api/statistics/managers', methods=['GET'])
    def get_statistics_managers():
        date_from_str = request.args.get('date_from', '')
        date_to_str = request.args.get('date_to', '')
        
        try:
            if date_from_str:
                date_from = datetime.fromisoformat(date_from_str.replace('T', ' '))
            else:
                date_from = datetime.utcnow().replace(day=1, hour=0, minute=0, second=0, microsecond=0)
            
            if date_to_str:
                date_to = datetime.fromisoformat(date_to_str.replace('T', ' '))
            else:
                date_to = datetime.utcnow()
        except:
            return jsonify({'error': 'Неверный формат даты'}), 400
        
        date_to = date_to.replace(hour=23, minute=59, second=59)
        
        with get_db() as db:
            # Считаем выполненные задачи по каждому менеджеру
            orders = db.query(Order).filter(
                Order.status == 'completed',
                Order.completed_at >= date_from,
                Order.completed_at <= date_to,
                Order.completed_by.isnot(None)
            ).all()
            
            personal_tasks = db.query(PersonalTask).filter(
                PersonalTask.status == 'completed',
                PersonalTask.completed_at >= date_from,
                PersonalTask.completed_at <= date_to,
                PersonalTask.completed_by.isnot(None)
            ).all()
            
            # Группируем по completed_by
            managers = {}
            
            for order in orders:
                name = order.completed_by
                if name:
                    if name not in managers:
                        managers[name] = {'name': name, 'orders_count': 0, 'personal_tasks_count': 0, 'total': 0}
                    managers[name]['orders_count'] += 1
                    managers[name]['total'] += 1
            
            for task in personal_tasks:
                name = task.completed_by
                if name:
                    if name not in managers:
                        managers[name] = {'name': name, 'orders_count': 0, 'personal_tasks_count': 0, 'total': 0}
                    managers[name]['personal_tasks_count'] += 1
                    managers[name]['total'] += 1
            
            # Сортируем по убыванию
            result = sorted(managers.values(), key=lambda x: x['total'], reverse=True)
            
            return jsonify(result), 200
    
    # ===== GET: Экспорт в Excel =====
    @app.route('/api/statistics/export', methods=['GET'])
    def export_statistics():
        if not OPENPYXL_AVAILABLE:
            return jsonify({'error': 'openpyxl не установлен'}), 500
        
        date_from_str = request.args.get('date_from', '')
        date_to_str = request.args.get('date_to', '')
        
        try:
            if date_from_str:
                date_from = datetime.fromisoformat(date_from_str.replace('T', ' '))
            else:
                date_from = datetime.utcnow().replace(day=1, hour=0, minute=0, second=0, microsecond=0)
            
            if date_to_str:
                date_to = datetime.fromisoformat(date_to_str.replace('T', ' '))
            else:
                date_to = datetime.utcnow()
        except:
            return jsonify({'error': 'Неверный формат даты'}), 400
        
        date_to = date_to.replace(hour=23, minute=59, second=59)
        
        # Создаём Excel
        wb = Workbook()
        
        # ===== Лист 1: Показатели =====
        ws1 = wb.active
        ws1.title = 'Показатели'
        
        # Стили
        header_font = Font(bold=True, size=12)
        header_fill = PatternFill(start_color='667EEA', end_color='667EEA', fill_type='solid')
        header_font_white = Font(bold=True, color='FFFFFF', size=12)
        
        ws1['A1'] = 'Показатель'
        ws1['B1'] = 'Значение'
        ws1['A1'].font = header_font_white
        ws1['B1'].font = header_font_white
        ws1['A1'].fill = header_fill
        ws1['B1'].fill = header_fill
        
        with get_db() as db:
            # Собираем данные
            shipments_completed = db.query(Order).filter(
                Order.type.in_(['regions', 'spb']),
                Order.status == 'completed',
                Order.completed_at >= date_from,
                Order.completed_at <= date_to
            ).count()
            
            shipments_with_errors = db.query(OrderError).join(
                Order, Order.id == OrderError.order_id
            ).filter(
                Order.type.in_(['regions', 'spb']),
                OrderError.created_at >= date_from,
                OrderError.created_at <= date_to
            ).distinct(OrderError.order_id).count()
            
            orders_created = db.query(Order).filter(
                Order.created_at >= date_from,
                Order.created_at <= date_to
            ).count()
            
            personal_tasks_created = db.query(PersonalTask).filter(
                PersonalTask.created_at >= date_from,
                PersonalTask.created_at <= date_to
            ).count()
            
            orders_completed = db.query(Order).filter(
                Order.status == 'completed',
                Order.completed_at >= date_from,
                Order.completed_at <= date_to
            ).count()
            
            personal_tasks_completed = db.query(PersonalTask).filter(
                PersonalTask.status == 'completed',
                PersonalTask.completed_at >= date_from,
                PersonalTask.completed_at <= date_to
            ).count()
            
            complaints_confirmed = db.query(Complaint).filter(
                Complaint.status == 'confirmed',
                Complaint.created_at >= date_from,
                Complaint.created_at <= date_to
            ).count()
            
            # Записываем данные
            ws1['A3'] = 'Период'
            ws1['B3'] = f'{date_from.strftime("%d.%m.%Y")} — {date_to.strftime("%d.%m.%Y")}'
            
            ws1['A5'] = 'Отгрузок выполнено'
            ws1['B5'] = shipments_completed
            
            ws1['A6'] = 'Отгрузок с ошибками'
            ws1['B6'] = shipments_with_errors
            
            ws1['A7'] = 'Задач поставлено'
            ws1['B7'] = orders_created + personal_tasks_created
            
            ws1['A8'] = 'Задач выполнено'
            ws1['B8'] = orders_completed + personal_tasks_completed
            
            ws1['A9'] = 'Жалоб подтверждено'
            ws1['B9'] = complaints_confirmed
            
            # Детализация
            ws1['A11'] = 'Детализация'
            ws1['A11'].font = header_font
            ws1['A12'] = 'Задач из хабов создано'
            ws1['B12'] = orders_created
            ws1['A13'] = 'Личных задач создано'
            ws1['B13'] = personal_tasks_created
            ws1['A14'] = 'Задач из хабов выполнено'
            ws1['B14'] = orders_completed
            ws1['A15'] = 'Личных задач выполнено'
            ws1['B15'] = personal_tasks_completed
            
            # ===== Лист 2: Ошибки =====
            ws2 = wb.create_sheet('Ошибки')
            ws2['A1'] = 'Номер заказа'
            ws2['B1'] = 'Контрагент'
            ws2['C1'] = 'Тип ошибки'
            ws2['D1'] = 'Описание'
            ws2['E1'] = 'Автор'
            ws2['F1'] = 'Дата'
            
            for col in ['A1', 'B1', 'C1', 'D1', 'E1', 'F1']:
                ws2[col].font = header_font_white
                ws2[col].fill = header_fill
            
            errors = db.query(OrderError).filter(
                OrderError.created_at >= date_from,
                OrderError.created_at <= date_to
            ).order_by(OrderError.created_at.desc()).all()
            
            row = 2
            for error in errors:
                error_type = db.query(ErrorType).filter(ErrorType.id == error.error_type_id).first()
                order = db.query(Order).filter(Order.id == error.order_id).first()
                
                email_data = {}
                if order and order.email_data:
                    try:
                        email_data = json.loads(order.email_data)
                    except:
                        pass
                
                ws2[f'A{row}'] = email_data.get('order_number', '')
                ws2[f'B{row}'] = email_data.get('contractor', '')
                ws2[f'C{row}'] = error_type.name if error_type else ''
                ws2[f'D{row}'] = error.description or ''
                ws2[f'E{row}'] = error.author
                ws2[f'F{row}'] = error.created_at.strftime('%d.%m.%Y %H:%M') if error.created_at else ''
                row += 1
            
            # ===== Лист 3: Жалобы =====
            ws3 = wb.create_sheet('Жалобы')
            ws3['A1'] = 'Имя'
            ws3['B1'] = 'Категория'
            ws3['C1'] = 'Текст'
            ws3['D1'] = 'Статус'
            ws3['E1'] = 'Дата'
            
            for col in ['A1', 'B1', 'C1', 'D1', 'E1']:
                ws3[col].font = header_font_white
                ws3[col].fill = header_fill
            
            complaints = db.query(Complaint).filter(
                Complaint.created_at >= date_from,
                Complaint.created_at <= date_to
            ).order_by(Complaint.created_at.desc()).all()
            
            row = 2
            for complaint in complaints:
                category = db.query(ComplaintCategory).filter(
                    ComplaintCategory.id == complaint.category_id
                ).first()
                
                ws3[f'A{row}'] = complaint.name
                ws3[f'B{row}'] = category.name if category else ''
                ws3[f'C{row}'] = complaint.text
                ws3[f'D{row}'] = complaint.status
                ws3[f'E{row}'] = complaint.created_at.strftime('%d.%m.%Y %H:%M') if complaint.created_at else ''
                row += 1
            
            # ===== Лист 4: Загрузка менеджеров =====
            ws4 = wb.create_sheet('Менеджеры')
            ws4['A1'] = 'Менеджер'
            ws4['B1'] = 'Задач из хабов'
            ws4['C1'] = 'Личных задач'
            ws4['D1'] = 'Всего'
            
            for col in ['A1', 'B1', 'C1', 'D1']:
                ws4[col].font = header_font_white
                ws4[col].fill = header_fill
            
            # Собираем данные по менеджерам
            managers = {}
            
            orders_list = db.query(Order).filter(
                Order.status == 'completed',
                Order.completed_at >= date_from,
                Order.completed_at <= date_to,
                Order.completed_by.isnot(None)
            ).all()
            
            for order in orders_list:
                name = order.completed_by
                if name not in managers:
                    managers[name] = {'orders': 0, 'personal': 0, 'total': 0}
                managers[name]['orders'] += 1
                managers[name]['total'] += 1
            
            personal_list = db.query(PersonalTask).filter(
                PersonalTask.status == 'completed',
                PersonalTask.completed_at >= date_from,
                PersonalTask.completed_at <= date_to,
                PersonalTask.completed_by.isnot(None)
            ).all()
            
            for task in personal_list:
                name = task.completed_by
                if name not in managers:
                    managers[name] = {'orders': 0, 'personal': 0, 'total': 0}
                managers[name]['personal'] += 1
                managers[name]['total'] += 1
            
            # Сортируем
            sorted_managers = sorted(managers.items(), key=lambda x: x[1]['total'], reverse=True)
            
            row = 2
            for name, data in sorted_managers:
                ws4[f'A{row}'] = name
                ws4[f'B{row}'] = data['orders']
                ws4[f'C{row}'] = data['personal']
                ws4[f'D{row}'] = data['total']
                row += 1
        
        # Сохраняем в буфер
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        # Формируем имя файла
        filename = f'statistics_{date_from.strftime("%Y%m%d")}_{date_to.strftime("%Y%m%d")}.xlsx'
        
        return send_file(
            output,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name=filename
        )